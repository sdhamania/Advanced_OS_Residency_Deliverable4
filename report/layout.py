#!/usr/bin/env python3
"""Shared APA 7 layout helpers for this deliverable."""

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "report" / "Advanced_Shell_Deliverable_4_Report.docx"
TITLE = "Advanced Shell Simulation with Integrated OS Concepts"
SUBTITLE = "Deliverable 4 Integration and Security Implementation"


def font(run, *, bold=False, italic=False, size=12):
    run.font.name = "Times New Roman"
    properties = run._element.get_or_add_rPr()
    fonts = properties.rFonts
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + attribute), "Times New Roman")
    for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn("w:" + attribute), None)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.bold = bold
    run.italic = italic
    return run


def paragraph(doc, text="", *, indent=True, bold=False, italic=False,
              centered=False, keep=False, style=None):
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.first_line_indent = Inches(0.5 if indent else 0)
    pf.left_indent = Inches(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 2
    pf.widow_control = True
    pf.keep_with_next = keep
    font(p.add_run(text), bold=bold, italic=italic)
    return p


def heading(doc, text, level=1):
    return paragraph(doc, text, indent=False, bold=True, centered=level == 1,
                     keep=True, style=f"Heading {level}")


def page(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1
    p.add_run().add_break(__import__("docx").enum.text.WD_BREAK.PAGE)


def caption(doc, label, title):
    paragraph(doc, label, indent=False, bold=True, keep=True)
    paragraph(doc, title, indent=False, italic=True, keep=True)


def note(doc, text):
    p = paragraph(doc, indent=False)
    font(p.add_run("Note. "), italic=True)
    font(p.add_run(text))
    return p


def table(doc, label, title, headers, rows, widths):
    caption(doc, label, title)
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for column, width in zip(t.columns, widths):
        column.width = Inches(width)
    # APA 7 overrides the general document preset: horizontal rules only.
    props = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "bottom", "left", "right", "insideH", "insideV"):
        border = OxmlElement("w:" + side)
        border.set(qn("w:val"), "single" if side in ("top", "bottom") else "nil")
        border.set(qn("w:sz"), "6")
        border.set(qn("w:color"), "000000")
        borders.append(border)
    props.append(borders)
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row_index, values in enumerate([headers] + rows):
        row = t.rows[0] if row_index == 0 else t.add_row()
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        for index, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_props = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side, padding in (("top", 65), ("bottom", 65), ("left", 95), ("right", 95)):
                element = OxmlElement("w:" + side)
                element.set(qn("w:w"), str(padding))
                element.set(qn("w:type"), "dxa")
                margins.append(element)
            cell_props.append(margins)
            if row_index == 0:
                header_borders = OxmlElement("w:tcBorders")
                bottom = OxmlElement("w:bottom")
                bottom.set(qn("w:val"), "single")
                bottom.set(qn("w:sz"), "6")
                bottom.set(qn("w:color"), "000000")
                header_borders.append(bottom)
                cell_props.append(header_borders)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1
            font(p.add_run(value), bold=row_index == 0)
    return t


def figure(doc, number, title, filename, description, new_page=False):
    before = len(doc.paragraphs)
    caption(doc, f"Figure B{number}", title)
    doc.paragraphs[before].paragraph_format.page_break_before = new_page
    p = paragraph(doc, indent=False)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.keep_with_next = True
    from PIL import Image
    path = ROOT / "screenshots" / filename
    with Image.open(path) as source:
        width, height = source.size
    display_width = min(6.5, 5.5 * width / height)
    picture = p.add_run().add_picture(str(path), width=Inches(display_width))
    picture._inline.docPr.set("descr", description)
    note(doc, description)


def configure(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, side, Inches(1))
    section.header_distance = Inches(0.5)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Caption"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.line_spacing = 2
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        properties = style._element.get_or_add_pPr()
        border = properties.find(qn("w:pBdr"))
        if border is not None:
            properties.remove(border)
    p = section.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.line_spacing = 1
    r = p.add_run()
    font(r)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    field_run = OxmlElement("w:r")
    field_text = OxmlElement("w:t")
    field_text.text = "1"
    field_run.append(field_text)
    field.append(field_run)
    p._p.append(field)
    doc.core_properties.author = "Shobhit Dhamania"
    doc.core_properties.title = TITLE
    doc.core_properties.subject = SUBTITLE

