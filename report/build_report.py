#!/usr/bin/env python3
"""Build the D4 APA report; append only original native screenshots."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
import layout as L
ROOT=Path(__file__).resolve().parents[1]
OUTPUT=ROOT/'report/Advanced_Shell_Deliverable_4_Report.docx'
TITLE='Advanced Shell Simulation with Integrated OS Concepts'
SUBTITLE='Deliverable 4 Integration and Security Implementation'
FIGURES=[
('01_admin_login.png','Administrator Authentication and a Permitted System Update','The administrator logs in and appends a line to the simulated system configuration.'),
('02_pipelines.png','Piping Built In Output and Chaining Three Commands','The shell filters a directory listing and sorts the error lines from a file through three concurrent stages.'),
('03_permissions_admin.png','Administrator Control of Resource Permissions','Mode 604 gives other users read permission. The administrator has an explicit override.'),
('04_student_permissions.png','Standard User Access to Protected System Files','The student can read the system configuration but cannot append to or remove it. Read is allowed for task.txt; write and execute are denied.'),
('05_user_ownership.png','Private File Ownership and an Approved Pipeline','A newly created student file has owner student and mode 600. A host path is denied while the approved streaming pipeline succeeds.'),
('06_guest_denial.png','A Different Standard User Cannot Read the Student Home','The guest authenticates with its own credentials and is denied access through the student home directory.'),
('07_execute_allowed.png','Execute Permission Granted Without Write Permission','After the administrator changes task.txt to mode 605, the student execute check is allowed and the write check remains denied. This is a simulated access decision.'),
('08_failed_login.png','Invalid Credentials Prevent Shell Access','Three incorrect passwords end the session without exposing a shell command prompt.')]

def p(doc,text):return L.paragraph(doc,text)
def code(doc,text):
    for line in text.splitlines():
        para=L.paragraph(doc,indent=False)
        para.paragraph_format.line_spacing=1
        para.paragraph_format.keep_with_next=True
        run=para.add_run(line);run.font.name='Courier New';run.font.size=Pt(10)
    doc.paragraphs[-1].paragraph_format.keep_with_next=False

def main():
    data=json.loads((ROOT/'evidence/results.json').read_text())
    assert data['verification']['integration_tests']==71
    present=[(ROOT/'screenshots'/name).is_file() for name,_,_ in FIGURES]
    if any(present) and not all(present):
        raise SystemExit('Incomplete native screenshot set: '+', '.join(name for (name,_,_),found in zip(FIGURES,present) if not found))
    doc=Document(ROOT/'report/apa_template.docx')
    title=L.paragraph(doc,TITLE+'\n'+SUBTITLE,indent=False,centered=True,bold=True,style='Title')
    title.paragraph_format.space_before=Pt(72)
    L.paragraph(doc,indent=False)
    for line in ['Shobhit Dhamania','University of the Cumberlands','MSCS-630-M20: Advanced Operating Systems','Sean Perryman','October 25, 2026']:
        L.paragraph(doc,line,indent=False,centered=True)
    L.page(doc)
    L.heading(doc,TITLE+'\n'+SUBTITLE)
    p(doc,'Deliverable 4 combines the earlier shell components with command pipelines, authenticated sessions, and simulated file authorization. My main design decision was to keep the existing module interfaces and add checks where commands enter the execution path. This allows a pipeline such as cat shared/events.txt | grep error | sort to coexist with Round Robin scheduling, priority scheduling, paging, and the producer-consumer demonstration in one shell. The completed implementation passes 71 integration tests and the parser and scheduler unit suites on macOS.')
    p(doc,'The most useful insight came from considering what a pipe changes about security. A file-reading command at the beginning of a pipeline still needs authorization, even though the final command only sees a stream. Checking every stage before creating any child prevents an unauthorized stage from being launched after an earlier stage has already started. That observation helped connect process creation, descriptor ownership, and access control into one design.')
    L.heading(doc,'Integration Through Existing Interfaces')
    p(doc,'The D1 shell remains responsible for parsing, built-ins, process groups, terminal ownership, and jobs. D2 retains Scheduler::command, including its configurable Round Robin quantum and priority heap with preemption at simulated arrival events. D3 retains PagingMemory::command and synchronizationCommand, including FIFO/LRU demand paging and real producer-consumer worker processes. Seven files containing those modules, their public interfaces, and the Job model are byte-identical to D3; evidence/preserved_modules.json records their SHA-256 hashes.')
    p(doc,'The parser now attaches a vector of argument vectors to ParsedCommand for a pipeline. Security::authorize checks the entire command or every pipeline stage. A standalone foreground built-in then runs in the parent, preserving its state. Other commands use the existing launcher. For a pipeline, its child becomes a coordinator that calls executePipeline with a stage callback; the callback dispatches a built-in or executes an external binary. The new executor therefore does not need to know how scheduling or paging works.')
    L.table(doc,'Table 1','Responsibilities and State Boundaries',['Component','Interface','Responsibility'],[
        ['D1 shell','run / launchExternal','Authentication entry, dispatch, process group and terminal control'],
        ['D2 scheduler','Scheduler::command','Logical CPU schedule and per-job timing metrics'],
        ['D3 paging','command(args)','Logical pages, frames, replacement and fault counters'],
        ['D3 sync','sync(args)','Forked workers, shared ring and named semaphores'],
        ['D4 security','authorize / command','Identity, ownership, persistent modes and access decisions'],
        ['D4 pipelines','executePipeline','Concurrent stages, descriptor wiring and reaping']],[1.2,1.8,3.5])
    L.note(doc,'Logical scheduling jobs and paging process IDs remain simulation objects. Pipeline stages and synchronization workers are actual host processes.')
    L.heading(doc,'Piping and Job Control')
    p(doc,'For a pipeline with N stages, the executor creates N minus one anonymous pipes. Each stage is forked before the coordinator waits for any of them. The first stage inherits standard input, intermediate stages connect to both adjacent pipes, and the final stage inherits standard output. dup2 maps the appropriate pipe endpoints onto descriptors 0 and 1. POSIX specifies that dup2 replaces the destination descriptor when necessary (IEEE & The Open Group, 2024a). Error output remains on the terminal, so diagnostic messages do not become input data for the next command.')
    p(doc,'Every stage closes all original pipe endpoints after duplication, and the coordinator closes its copies after launching the stages. This ownership rule is essential: a consumer cannot observe end-of-file while an unused writer remains open. Starting stages concurrently also prevents a producer from filling a bounded pipe while the shell waits for it to finish before launching the consumer. The large-transfer test sends 700,000 bytes through cat | cat | wc -c and verifies the exact byte count.')
    p(doc,'One D1 process group contains the coordinator and all stages. Foreground terminal assignment, Ctrl+Z, bg, fg, and Ctrl+C apply to the whole pipeline. The coordinator ignores interrupt, quit, hangup, and termination signals while stages use default dispositions, so it can reap them. Its normal stop disposition exposes group suspension to the shell. The kill built-in maps a tracked job PID to its group and continues stopped jobs to deliver termination. Completion returns the last stage status; failures in earlier stages do not enable pipefail.')
    p(doc,'The grammar permits two through sixteen stages and a final background marker. Quoted or escaped pipes remain literal arguments. Empty stages and unsupported redirection or command lists produce errors. cd and exit inside a pipeline affect only their child process. Scheduling, paging, synchronization, file mutations, permission changes, and fg/bg require standalone foreground execution, because a child would otherwise change a disposable copy of parent state or compete to persist metadata.')
    L.heading(doc,'Authentication and Session Identity')
    p(doc,'Before the command loop begins, the shell asks for a username and password. Provisioning creates separate admin, student, and guest accounts; the latter two have standard roles. There are no shipped default passwords or command-line password options. Password input disables terminal echo before displaying its prompt and restores the previous mode afterward. Three failed attempts end the process. Unknown usernames receive the same failure response and perform the same password-derivation work as known accounts.')
    p(doc,'The Python provisioning utility creates a distinct random 16-byte salt for each account and a 32-byte PBKDF2-HMAC-SHA256 verifier using 600,000 iterations. C++ verifies those records through CommonCrypto on macOS or OpenSSL on Linux and compares every derived byte. OWASP documents this PBKDF2 work factor while preferring memory-hard algorithms for new production systems (OWASP Foundation, n.d.). I selected the platform implementations to avoid writing cryptographic primitives in the project. This choice does not establish FIPS certification or make the educational shell a production authentication service.')
    p(doc,'Credentials and the simulated permission database are stored in a host-private state directory outside the exposed workspace. The shell checks ownership and rejects group or other access to these files. It also obtains an exclusive session lock before loading them, preventing two shells from reading and later overwriting conflicting permission snapshots. Account changes use a fresh authenticated launch. Simulation state ends with that session, while file ownership and modes persist.')
    L.heading(doc,'File Permissions and Role Restrictions')
    p(doc,'The permission model stores an owner and three octal digits for each resource. Read, write, and execute/search use values 4, 2, and 1. Owner bits apply to the owner, and other bits apply to everyone else. Group bits are retained in the displayed mode but have no effect because this model has no simulated groups. POSIX defines distinct owner, group, and other permission classes, so this is a deliberate simplification of that model (IEEE & The Open Group, 2024b).')
    p(doc,'File reads require read permission, appending or touching an existing file requires write permission, and directory traversal requires search permission along every ancestor. Creating or removing an entry requires write and search on its parent directory. A newly created file belongs to the current user with mode 600; a directory begins at 700. Owners can change their modes, and only admin can transfer ownership. The access command makes read, write, and execute decisions visible. An execute check evaluates permission for an existing resource; it does not launch arbitrary standard-user programs.')
    p(doc,'Standard users begin in their own homes and are confined to a canonical workspace path. Canonicalization resolves symbolic links before authorization, preventing a link inside the workspace from pointing a file read outside it. Unknown resources have no default permission. The simulated system directory also has an explicit role rule: a standard user cannot modify or remove its contents even if an administrator grants permissive mode bits. A default mode of 644 allows the system configuration to be read while preserving administrator control of updates.')
    p(doc,'Allowing unrestricted external commands would undermine these checks because a standard user could launch another shell and read files directly under the same host UID. Standard users therefore receive a small external interface: grep with a pattern and optional files, sort and wc with optional files, and numeric sleep. File arguments are authorized, option arguments are rejected, and fixed binary paths avoid substitution through PATH. Host-PID kill and arbitrary external execution remain administrator capabilities. Pipelines use this same policy at every stage before any process is created.')
    L.heading(doc,'Verification and Performance Analysis')
    p(doc,'The 71 integration tests comprise 20 D1 process and terminal tests, four D2 scheduler tests, 20 D3 memory and synchronization tests, and 27 D4 tests. The D2 oracle checks 40 workloads under both policies. D4 covers failed and retried login, distinct identities, hidden passwords, persistent permissions, private homes, path and symlink escapes, external-command bypass attempts, session locking, and pipeline job control. A regression test kills running and stopped pipelines and verifies that all stages exit. All 27 D4 tests pass under AddressSanitizer and UndefinedBehaviorSanitizer.')
    L.table(doc,'Table 2','Observed Integration Results',['Experiment','Observed result','Interpretation'],[
        ['Three-stage filter','error alpha before error zebra','Output crosses both pipes and is sorted'],
        ['Large transfer','700,000 bytes received','Concurrent stages complete beyond pipe capacity'],
        ['Standard system update','Append and remove denied','Read access does not imply write authority'],
        ['Round Robin q=2','W=2.500 T=5.000 R=1.500','D2 reference workload remains consistent'],
        ['Priority scheduling','W=2.750 T=5.250 R=2.000','Higher-priority arrival preempts A at time 2'],
        ['FIFO / LRU','7 / 5 faults for 11 references','D3 replacement behavior remains consistent'],
        ['Producer-consumer stress','8,000 unique items; 0 errors','Exact-once transfer through a seven-slot ring']],[1.5,2.2,2.8])
    L.note(doc,'W is average waiting time, T is average turnaround time, and R is average response time in simulated units. Scheduling uses A=(6,3,0), B=(2,1,2), C=(1,1,2), D=(1,3,3), expressed as burst, priority, arrival. Lower priority numbers rank higher.')
    p(doc,'The evidence collector starts seven isolated authenticated sessions and saves actual output without passwords or verifier records. In this run, the pipeline session took '+str(data['pipelines']['session_wall_ms'])+' ms and the large-transfer session took '+str(data['large_pipeline']['session_wall_ms'])+' ms. These measurements include password derivation, shell startup, and captured output; they are not estimates of CPU scheduling overhead or isolated pipeline throughput. The failed-login session took '+str(data['failed_login']['session_wall_ms'])+' ms across three derivations. A meaningful throughput comparison would time only command execution after login and repeat trials under controlled load.')
    L.heading(doc,'Challenges and Improvements')
    p(doc,'Integration exposed two different kinds of state. Scheduling and paging should remain in the shell process, while a pipeline should run independent stages concurrently. I addressed that conflict by rejecting stateful module commands in pipelines and background execution rather than letting them silently update a child copy. The same rule protects permission metadata and the file operations that maintain it. This keeps the public interfaces understandable and makes the limitation visible to the user.')
    p(doc,'The descriptor design was another important connection. A command can produce correct output and still leave its consumer waiting forever if a writer descriptor remains open elsewhere. Treating pipe closure as part of process ownership made the larger transfer and early-consumer tests useful checks rather than relying only on a short echo demonstration. The terminal tests then connected that resource ownership to D1 process groups: a pipeline should stop and resume as one job.')
    p(doc,'The security boundary is intentionally narrower than a real OS boundary. All roles share the launching host UID, admin is trusted, and the host account owner can edit the program or state directly. Canonical-path checks do not defend against a hostile host process changing files between a check and an operation. The application does not provide MFA, persistent brute-force lockout, password recovery, group membership, or privilege separation. Production improvements would include OS-enforced identities or a sandbox, descriptor-relative file operations, a memory-hard password scheme, and transactional permission storage. Linux support is implemented but this submission was executed on macOS.')
    L.heading(doc,'References').paragraph_format.page_break_before=True
    refs=[
        ('IEEE & The Open Group. (2024a). ', 'dup, dup2, dup3', '. The Open Group Base Specifications Issue 8. https://pubs.opengroup.org/onlinepubs/9799919799/functions/dup.html'),
        ('IEEE & The Open Group. (2024b). ', 'General concepts', '. The Open Group Base Specifications Issue 8. https://pubs.opengroup.org/onlinepubs/9799919799/basedefs/V1_chap04.html'),
        ('OWASP Foundation. (n.d.). ', 'Password storage cheat sheet', '. OWASP Cheat Sheet Series. Retrieved October 25, 2026, from https://cheatsheetseries.owasp.org/cheatsheets/Password_Storage_Cheat_Sheet.html')]
    for author,title,rest in refs:
        para=L.paragraph(doc,indent=False);para.paragraph_format.left_indent=Inches(.5);para.paragraph_format.first_line_indent=Inches(-.5)
        L.font(para.add_run(author));L.font(para.add_run(title),italic=True);L.font(para.add_run(rest))
    L.heading(doc,'Appendix A').paragraph_format.page_break_before=True;L.heading(doc,'Reproducing the Integrated Demonstration')
    p(doc,'From the repository, run make, then python3 scripts/setup_security.py and choose distinct disposable passwords. Launch ./build/os-shell and log in as admin. From the repository launch directory, enter the following commands one at a time. The collector reproduces the same workloads automatically in a temporary state.')
    code(doc,'cd .osh-security/workspace\nls shared | grep txt\ncat shared/events.txt | grep error | sort\nwrite system/config.txt admin-update\nchmod 604 shared/task.txt\nexit 0')
    p(doc,'Launch the shell again and authenticate as student. Its initial directory is home/student. The following operations should show a readable system file, a denied write, and explicit read/write/execute decisions. Use screenshots/CAPTURE_GUIDE.md for the complete native capture sequence, including guest access and a later execute grant. Figures B1 through B8 present the original native Terminal captures of these demonstrations.')
    code(doc,'cat ../../system/config.txt\nwrite ../../system/config.txt forbidden\naccess r ../../shared/task.txt\naccess w ../../shared/task.txt\naccess x ../../shared/task.txt\ncat ../../shared/events.txt | grep error | sort\nexit 0')
    if all(present):
        for i,(name,title,description) in enumerate(FIGURES,1):
            if i==1:
                L.heading(doc,'Appendix B').paragraph_format.page_break_before=True
                L.heading(doc,'Native Terminal Evidence')
            L.figure(doc,i,title,name,description,new_page=i>1)
    doc.core_properties.title=TITLE+' '+SUBTITLE
    doc.core_properties.subject='Integration, piping, authentication, and simulated file permissions'
    doc.save(OUTPUT)
    print(OUTPUT)
if __name__=='__main__':main()
